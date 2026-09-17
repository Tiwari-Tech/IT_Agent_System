from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader


def extract_text(path: Path, file_type: str) -> tuple[str, dict]:
    file_type = file_type.lower().lstrip(".")
    if file_type == "pdf":
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages), {"pages": len(pages)}
    if file_type == "docx":
        doc = DocxDocument(str(path))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs), {"paragraphs": len(doc.paragraphs)}
    if file_type == "xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        rows: list[str] = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    rows.append(" | ".join(values))
        return "\n".join(rows), {"sheets": len(wb.worksheets)}
    if file_type in {"html", "htm"}:
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
        return soup.get_text("\n"), {}
    if file_type == "txt":
        return path.read_text(encoding="utf-8", errors="ignore"), {}
    raise ValueError("unsupported document type")
